#!/usr/bin/env python3
"""Generate style-driven CUMCM submission and optional print DOCX files."""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


HEADING_RE = re.compile(r"^(?P<marks>#{1,6})\s+(?P<title>.+?)\s*$")
IMAGE_RE = re.compile(r"^!\[(?P<caption>.*?)\]\((?P<path>[^)]+)\)\s*$")
KEYWORDS_RE = re.compile(r"^\s*(?:\*\*)?关键词(?:\*\*)?\s*[：:]\s*(?P<body>.+)$")
TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?(?:\s*:?-+:?\s*\|)+\s*$")
NUMBERED_HEADING_RE = re.compile(r"^(?:第?[一二三四五六七八九十百]+[、.]|\d+(?:\.\d+)*[、.]?)\s*")
BEGIN_EQUATION_RE = re.compile(r"\\begin\{equation(?P<star>\*)?\}")
END_EQUATION_RE = re.compile(r"\\end\{equation\*?\}")
DISPLAY_OPEN_RE = re.compile(r"(?<!\\)\\\[")
DISPLAY_CLOSE_RE = re.compile(r"(?<!\\)\\\]")


@dataclass(frozen=True)
class StyleSpec:
    east_asia: str
    latin: str
    size: float
    bold: bool
    alignment: int
    line: float
    before: float
    after: float
    first_chars: int = 0
    hanging_chars: int = 0
    keep_next: bool = False
    keep_lines: bool = False
    widow: bool = True


STYLE_SPECS = {
    "PaperTitle": StyleSpec("黑体", "Times New Roman", 18, True, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 0, 12, keep_next=True),
    "AbstractHeading": StyleSpec("黑体", "Times New Roman", 16, True, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 0, 6, keep_next=True),
    "AbstractBody": StyleSpec("宋体", "Times New Roman", 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 0, 0, first_chars=200),
    "Keywords": StyleSpec("宋体", "Times New Roman", 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 6, 0),
    "Heading1": StyleSpec("黑体", "Times New Roman", 16, True, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 12, 6, keep_next=True, keep_lines=True),
    "Heading2": StyleSpec("黑体", "Times New Roman", 14, True, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 8, 4, keep_next=True, keep_lines=True),
    "Heading3": StyleSpec("黑体", "Times New Roman", 12, True, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 6, 3, keep_next=True, keep_lines=True),
    "BodyText": StyleSpec("宋体", "Times New Roman", 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 0, 0, first_chars=200),
    "ListText": StyleSpec("宋体", "Times New Roman", 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 0, 0, hanging_chars=200),
    "FigureCaption": StyleSpec("宋体", "Times New Roman", 10.5, False, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 3, 6),
    "TableCaption": StyleSpec("宋体", "Times New Roman", 10.5, False, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 6, 3, keep_next=True),
    "TableText": StyleSpec("宋体", "Times New Roman", 10.5, False, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 0, 0),
    "EquationBlock": StyleSpec("Cambria Math", "Times New Roman", 12, False, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 6, 6, keep_lines=True),
    "Reference": StyleSpec("宋体", "Times New Roman", 10.5, False, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0, 0, hanging_chars=200, keep_lines=True),
    "AppendixHeading": StyleSpec("黑体", "Times New Roman", 16, True, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 12, 6, keep_next=True, keep_lines=True),
    "CodeBlock": StyleSpec("等线", "Consolas", 8.5, False, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0, 0),
}


def _set_style_fonts(style, spec: StyleSpec) -> None:
    style.font.name = spec.latin
    style.font.size = Pt(spec.size)
    style.font.bold = spec.bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), spec.latin)
    rfonts.set(qn("w:hAnsi"), spec.latin)
    rfonts.set(qn("w:eastAsia"), spec.east_asia)
    rfonts.set(qn("w:cs"), spec.latin)


def _set_char_indent(paragraph_format, first_chars: int = 0, hanging_chars: int = 0) -> None:
    ppr = paragraph_format._element.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:firstLineChars"), str(first_chars))
    ind.set(qn("w:hangingChars"), str(hanging_chars))
    ind.set(qn("w:left"), "0")
    ind.set(qn("w:right"), "0")


def _set_widow_control(paragraph_format, enabled: bool) -> None:
    ppr = paragraph_format._element.get_or_add_pPr()
    node = ppr.find(qn("w:widowControl"))
    if node is None:
        node = OxmlElement("w:widowControl")
        ppr.append(node)
    node.set(qn("w:val"), "1" if enabled else "0")


def ensure_styles(doc: DocumentType) -> None:
    for name, spec in STYLE_SPECS.items():
        builtin_name = {"Heading1": "Heading 1", "Heading2": "Heading 2", "Heading3": "Heading 3"}.get(name)
        if builtin_name:
            style = doc.styles[builtin_name]
            style.element.name.set(qn("w:val"), name)
        else:
            style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        _set_style_fonts(style, spec)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fmt = style.paragraph_format
        fmt.alignment = spec.alignment
        fmt.space_before = Pt(spec.before)
        fmt.space_after = Pt(spec.after)
        fmt.line_spacing = spec.line
        fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE if spec.line == 1.0 else WD_LINE_SPACING.MULTIPLE
        fmt.keep_with_next = spec.keep_next
        fmt.keep_together = spec.keep_lines
        _set_char_indent(fmt, spec.first_chars, spec.hanging_chars)
        _set_widow_control(fmt, spec.widow)
    _install_heading_numbering(doc)


def _install_heading_numbering(doc: DocumentType) -> None:
    numbering = doc.part.numbering_part.element
    existing_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(existing_ids, default=-1) + 1
    existing_num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing_num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)
    for level in range(3):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal")
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), ".".join(f"%{i + 1}" for i in range(level + 1)))
        suffix = OxmlElement("w:suff")
        suffix.set(qn("w:val"), "space")
        lvl.extend([start, num_fmt, lvl_text, suffix])
        abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)

    for level, style_name in enumerate(("Heading1", "Heading2", "Heading3")):
        style = doc.styles[style_name]
        ppr = style.element.get_or_add_pPr()
        old = ppr.find(qn("w:numPr"))
        if old is not None:
            ppr.remove(old)
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(level))
        num_id_node = OxmlElement("w:numId")
        num_id_node.set(qn("w:val"), str(num_id))
        num_pr.extend([ilvl, num_id_node])
        ppr.append(num_pr)


def configure_section(section, restart_page_number: bool = True) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.gutter = Cm(0)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    _clear_container(section.header)
    _set_page_footer(section, restart_page_number)


def _clear_container(container) -> None:
    for paragraph in container.paragraphs:
        for run in paragraph.runs:
            run._element.getparent().remove(run._element)


def _set_page_footer(section, restart: bool) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])
    if restart:
        sect_pr = section._sectPr
        pg_num = sect_pr.find(qn("w:pgNumType"))
        if pg_num is None:
            pg_num = OxmlElement("w:pgNumType")
            sect_pr.append(pg_num)
        pg_num.set(qn("w:start"), "1")


def _set_core_properties(doc: DocumentType) -> None:
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.title = ""
    props.subject = ""
    props.keywords = ""


def _add_text(paragraph, text: str, bold_prefix: str | None = None) -> None:
    if bold_prefix and text.startswith(bold_prefix):
        lead = paragraph.add_run(bold_prefix)
        lead.bold = True
        paragraph.add_run(text[len(bold_prefix) :])
    else:
        paragraph.add_run(text)


def _add_borderless_equation(doc: DocumentType, latex: str, number: int | None) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(1.5), Cm(13.0), Cm(1.5)]
    for cell, width in zip(table.rows[0].cells, widths):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.margin_top = 0
        cell.margin_bottom = 0
    table.rows[0].cells[0].paragraphs[0].style = "EquationBlock"
    center = table.rows[0].cells[1].paragraphs[0]
    center.style = "EquationBlock"
    center.alignment = WD_ALIGN_PARAGRAPH.CENTER
    center.add_run(latex)
    right = table.rows[0].cells[2].paragraphs[0]
    right.style = "EquationBlock"
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if number is not None:
        right.add_run(f"（{number}）")
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)


def _parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _set_cell_text(cell, text: str, header: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.style = "TableText"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = header
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _style_three_line_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, val, size in (("top", "single", "12"), ("bottom", "single", "12"), ("left", "nil", "0"), ("right", "nil", "0"), ("insideH", "nil", "0"), ("insideV", "nil", "0")):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), val)
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), "000000")
        borders.append(node)
    tbl_pr.append(borders)
    header_props = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_props.append(repeat)
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        cell_borders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:color"), "000000")
        cell_borders.append(bottom)
        tc_pr.append(cell_borders)


def _load_equation_mapping(mapping_path: Path | None) -> tuple[dict[int, int | None], dict[str, int]]:
    if not mapping_path:
        return {}, {}
    data = json.loads(mapping_path.read_text(encoding="utf-8"))
    by_line = {int(item["source_line_start"]): item.get("equation_number") for item in data.get("equations", []) if item.get("display")}
    by_label = {item["label"]: int(item["equation_number"]) for item in data.get("equations", []) if item.get("label") and item.get("equation_number") is not None}
    return by_line, by_label


def _display_latex(source: str) -> str:
    text = re.sub(r"\\begin\{equation\*?\}", "", source)
    text = re.sub(r"\\end\{equation\*?\}", "", text)
    text = re.sub(r"\\label\{[^{}]+\}", "", text)
    text = re.sub(r"^\s*\\\[", "", text)
    text = re.sub(r"\\\]\s*$", "", text)
    return "\n".join(line.strip() for line in text.strip().splitlines() if line.strip())


def _replace_equation_refs(text: str, labels: dict[str, int]) -> str:
    def eqref(match: re.Match[str]) -> str:
        label = match.group(1)
        return f"（{labels[label]}）" if label in labels else match.group(0)

    text = re.sub(r"\\eqref\{([^{}]+)\}", eqref, text)
    text = re.sub(r"\\ref\{([^{}]+)\}", lambda match: str(labels[match.group(1)]) if match.group(1) in labels else match.group(0), text)
    return text


def _suppress_paragraph_numbering(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "0")
    num_pr.append(num_id)
    ppr.append(num_pr)


def add_markdown_content(doc: DocumentType, paper_path: Path, mapping_path: Path | None = None) -> None:
    lines = paper_path.read_text(encoding="utf-8-sig").splitlines()
    equation_numbers, equation_labels = _load_equation_mapping(mapping_path)
    title_written = False
    abstract_seen = False
    body_started = False
    in_code = False
    code_buffer: list[str] = []
    figure_number = 0
    table_number = 0
    i = 0

    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()

        if stripped.startswith("```"):
            if in_code:
                paragraph = doc.add_paragraph(style="CodeBlock")
                paragraph.add_run("\n".join(code_buffer))
                code_buffer = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buffer.append(raw)
            i += 1
            continue
        if not stripped:
            i += 1
            continue

        heading = HEADING_RE.match(raw)
        if heading:
            title = heading.group("title").strip()
            level = len(heading.group("marks"))
            if not title_written:
                paragraph = doc.add_paragraph(style="PaperTitle")
                paragraph.add_run(title)
                title_written = True
            elif title == "摘要":
                paragraph = doc.add_paragraph(style="AbstractHeading")
                paragraph.add_run(title)
                abstract_seen = True
            else:
                if abstract_seen and not body_started:
                    doc.add_page_break()
                    body_started = True
                elif body_started and (title == "附录" or title.startswith("附录")):
                    doc.add_page_break()
                clean = NUMBERED_HEADING_RE.sub("", title)
                if title == "附录" or title.startswith("附录"):
                    style = "AppendixHeading"
                elif title == "参考文献":
                    style = "Heading1"
                else:
                    style = {1: "Heading1", 2: "Heading1", 3: "Heading2"}.get(level, "Heading3")
                paragraph = doc.add_paragraph(style=style)
                if title == "参考文献":
                    _suppress_paragraph_numbering(paragraph)
                paragraph.add_run(clean)
            i += 1
            continue

        keyword = KEYWORDS_RE.match(raw)
        if keyword:
            paragraph = doc.add_paragraph(style="Keywords")
            _add_text(paragraph, f"关键词：{keyword.group('body')}", "关键词：")
            i += 1
            continue

        image = IMAGE_RE.match(raw)
        if image:
            image_path = (paper_path.parent / image.group("path").strip()).resolve()
            if not image_path.exists():
                raise FileNotFoundError(f"image not found: {image_path}")
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(image_path), width=Cm(15.5))
            figure_number += 1
            caption = doc.add_paragraph(style="FigureCaption")
            caption.add_run(f"图{figure_number}  {image.group('caption')}")
            i += 1
            continue

        if "|" in raw and i + 1 < len(lines) and TABLE_SEPARATOR_RE.match(lines[i + 1]):
            rows = [_parse_table_row(raw)]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(_parse_table_row(lines[i]))
                i += 1
            table_number += 1
            caption = doc.add_paragraph(style="TableCaption")
            caption.add_run(f"表{table_number}  表格内容")
            columns = max(len(row) for row in rows)
            table = doc.add_table(rows=len(rows), cols=columns)
            width = Cm(16.0 / columns)
            for row_index, values in enumerate(rows):
                for col_index in range(columns):
                    cell = table.rows[row_index].cells[col_index]
                    cell.width = width
                    _set_cell_text(cell, values[col_index] if col_index < len(values) else "", header=row_index == 0)
            _style_three_line_table(table)
            continue

        begin = BEGIN_EQUATION_RE.search(raw)
        display = DISPLAY_OPEN_RE.search(raw)
        if begin or display:
            start_line = i + 1
            buffer = [raw]
            close_pattern = END_EQUATION_RE if begin else DISPLAY_CLOSE_RE
            while not close_pattern.search(buffer[-1]) and i + 1 < len(lines):
                i += 1
                buffer.append(lines[i])
            _add_borderless_equation(doc, _display_latex("\n".join(buffer).strip()), equation_numbers.get(start_line))
            i += 1
            continue

        style = "AbstractBody" if abstract_seen and not body_started else "BodyText"
        if stripped.startswith(('- ', '* ')) or re.match(r"^\d+[.)、]\s+", stripped):
            style = "ListText"
            stripped = re.sub(r"^(?:[-*]|\d+[.)、])\s+", "", stripped)
        if body_started and any(stripped.startswith(prefix) for prefix in ("[", "参考文献")) and re.match(r"^\[\d+\]", stripped):
            style = "Reference"
        paragraph = doc.add_paragraph(style=style)
        paragraph.add_run(_replace_equation_refs(stripped, equation_labels))
        i += 1


def build_document(paper: Path, output: Path, mapping: Path | None = None, official_frontmatter: Path | None = None) -> None:
    # This legacy renderer emits plain text, not OMML. Never silently downgrade math.
    from formula_compatibility_validator import extract_math
    if extract_math(paper.read_text(encoding='utf-8'))[0]:
        raise ValueError('WORD_NATIVE_EXPORT_PIPELINE = NOT_PRESENT in this legacy renderer; use a verified Pandoc/OMML exporter. Canonical Markdown is preserved.')
    doc = Document(str(official_frontmatter)) if official_frontmatter else Document()
    if official_frontmatter:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
    else:
        section = doc.sections[0]
    configure_section(section, restart_page_number=True)
    ensure_styles(doc)
    _set_core_properties(doc)
    add_markdown_content(doc, paper, mapping)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--paper", type=Path, required=True, help="Source paper.md")
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--equation-mapping", type=Path, help="equation_mapping.json")
    parser.add_argument("--official-frontmatter", type=Path, help="Authentic official commitment + numbering DOCX")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        submission = args.output_dir / "paper_submission.docx"
        build_document(args.paper, submission, args.equation_mapping)
        generated = [str(submission)]
        print_path = None
        if args.official_frontmatter:
            if not args.official_frontmatter.exists():
                raise FileNotFoundError(f"official front matter not found: {args.official_frontmatter}")
            print_path = args.output_dir / "paper_print.docx"
            build_document(args.paper, print_path, args.equation_mapping, args.official_frontmatter)
            generated.append(str(print_path))
        manifest = {
            "schema_version": "1.0",
            "source": str(args.paper),
            "generated": generated,
            "official_frontmatter": str(args.official_frontmatter) if args.official_frontmatter else None,
            "print_status": "generated-from-official-frontmatter" if print_path else "blocked-official-frontmatter-not-supplied",
            "equation_mode": "no-math-legacy-layout-only",
        }
        (args.output_dir / "word_generation_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"Word 文件已生成：{', '.join(generated)}")
        if not print_path:
            print("未生成纸质版：未提供真实官方前置页面，未使用程序仿制。")
        return 0
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        print(f"Word 生成失败：{exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())

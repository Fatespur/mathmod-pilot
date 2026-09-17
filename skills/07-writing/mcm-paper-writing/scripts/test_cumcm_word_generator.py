from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from cumcm_word_generator import STYLE_SPECS, build_document, main


PAPER = r"""# 示例论文标题

## 摘要

本文建立测试模型并验证文档样式。

关键词：测试；模型；验证

## 一、问题重述

这是正文段落，包含行内公式 \(x+y=1\)。

### 1.1 模型建立

\begin{equation}
\label{eq:test}
x+y=1
\end{equation}

## 参考文献

[1] 作者. 真实文献题名[J]. 期刊, 2026.

## 附录

```python
print("ok")
```
"""


MAPPING = {
    "equations": [
        {"source_line_start": 15, "equation_number": 1, "display": True}
    ]
}


class CumcmWordGeneratorTests(unittest.TestCase):
    def test_build_submission_styles_and_geometry(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            paper = root / "paper.md"
            mapping = root / "equation_mapping.json"
            output = root / "paper_submission.docx"
            paper.write_text(PAPER, encoding="utf-8")
            mapping.write_text(json.dumps(MAPPING), encoding="utf-8")
            build_document(paper, output, mapping)
            self.assertTrue(output.exists())
            doc = Document(output)
            self.assertTrue(set(STYLE_SPECS).issubset({style.name for style in doc.styles}))
            section = doc.sections[0]
            self.assertAlmostEqual(section.left_margin.cm, 2.5, places=1)
            self.assertAlmostEqual(section.page_width.cm, 21.0, places=1)
            self.assertEqual(doc.core_properties.author, "")
            self.assertTrue(any("PAGE" in paragraph._p.xml for paragraph in section.footer.paragraphs))

    def test_cli_does_not_fake_print_frontmatter(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            paper = root / "paper.md"
            paper.write_text(PAPER, encoding="utf-8")
            result = main(["--paper", str(paper), "--output-dir", str(root / "out")])
            self.assertEqual(result, 0)
            self.assertTrue((root / "out" / "paper_submission.docx").exists())
            self.assertFalse((root / "out" / "paper_print.docx").exists())
            manifest = json.loads((root / "out" / "word_generation_manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["print_status"], "blocked-official-frontmatter-not-supplied")

    def test_print_uses_supplied_frontmatter(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            paper = root / "paper.md"
            paper.write_text(PAPER, encoding="utf-8")
            front = root / "official.docx"
            official = Document()
            official.add_paragraph("官方固定承诺书内容")
            official.add_page_break()
            official.add_paragraph("官方固定编号页内容")
            official.save(front)
            result = main(["--paper", str(paper), "--output-dir", str(root / "out"), "--official-frontmatter", str(front)])
            self.assertEqual(result, 0)
            print_doc = Document(root / "out" / "paper_print.docx")
            text = "\n".join(item.text for item in print_doc.paragraphs)
            self.assertIn("官方固定承诺书内容", text)
            self.assertIn("示例论文标题", text)


if __name__ == "__main__":
    unittest.main()

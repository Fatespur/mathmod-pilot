from __future__ import annotations

import csv
import json
from pathlib import Path

from utils import optional_dependency_status, read_text


def parse_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("Reading PDF requires optional dependency pypdf") from exc
    return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)


def parse_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Reading DOCX requires optional dependency python-docx") from exc
    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        paragraphs.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
    return "\n".join(paragraphs)


def parse_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Reading XLSX requires optional dependency openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(max_row=30, values_only=True):
            lines.append("\t".join("" if value is None else str(value) for value in row))
    return "\n".join(lines)


def parse_csv(path: Path) -> str:
    lines = []
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        for index, row in enumerate(csv.reader(handle)):
            lines.append("\t".join(row))
            if index >= 30:
                break
    return "\n".join(lines)


def parse_document(path: Path) -> tuple[str, list[str]]:
    suffix = path.suffix.lower()
    warnings: list[str] = []
    try:
        if suffix == ".pdf":
            return parse_pdf(path), warnings
        if suffix == ".docx":
            return parse_docx(path), warnings
        if suffix == ".xlsx":
            return parse_xlsx(path), warnings
        if suffix == ".csv":
            return parse_csv(path), warnings
        if suffix == ".json":
            return json.dumps(json.loads(read_text(path)), ensure_ascii=False, indent=2), warnings
        return read_text(path), warnings
    except Exception as exc:
        warnings.append(f"{path.name}: {exc}")
        return "", warnings


def dependency_report() -> dict[str, bool]:
    return optional_dependency_status()
